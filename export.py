"""Отчёт в Excel из базы SQLite: профиль, меню на неделю, дневник веса с графиком, избранное.

Файл собирается в памяти и сразу уходит в чат — на сервере не хранится.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import storage
from nutrition import ACTIVITY, GOALS, SEX

GREEN = "2E7D32"
HEAD_FILL = PatternFill("solid", fgColor=GREEN)
HEAD_FONT = Font(bold=True, color="FFFFFF")
DAY_FILL = PatternFill("solid", fgColor="E8F5E9")
THIN = Border(bottom=Side(style="thin", color="C8E6C9"))


def _sheet(wb: Workbook, title: str, header: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def _profile(wb: Workbook, p) -> None:
    ws = _sheet(wb, "Профиль", ["Параметр", "Значение"], [34, 30])
    rows = [
        ("Имя", p["name"]), ("Пол", SEX[p["sex"]]), ("Возраст, лет", p["age"]),
        ("Вес по анкете, кг", p["weight"]), ("Рост, см", p["height"]),
        ("Активность", ACTIVITY[p["activity"]][0]), ("Цель", GOALS[p["goal"]][0]),
        ("Базовый обмен, ккал", p["bmr"]), ("С учётом активности, ккал", p["tdee"]),
        ("Норма на день, ккал", p["target"]),
        ("Белки, г", p["protein"]), ("Жиры, г", p["fat"]), ("Углеводы, г", p["carbs"]),
        ("Анкета обновлена (UTC)", p["updated_at"]),
    ]
    for r in rows:
        ws.append(r)
    ws.cell(row=11, column=1).font = ws.cell(row=11, column=2).font = Font(bold=True)


def _menu(wb: Workbook, menu: dict | None) -> None:
    ws = _sheet(wb, "Меню на неделю", ["День", "Приём пищи", "Блюдо", "Граммы", "Ккал"],
                [14, 14, 46, 10, 10])
    if not menu:
        ws.append(["Меню ещё не составлено — нажмите в боте «🍽 Меню на неделю»."])
        return
    for d in menu["days"]:
        ws.append([d["day"], "Итого за день", f"Б {d['protein']} · Ж {d['fat']} · У {d['carbs']} г",
                   None, d["kcal"]])
        for cell in ws[ws.max_row]:
            cell.fill, cell.font = DAY_FILL, Font(bold=True)
        for m in d["meals"]:
            for i in m["items"]:
                ws.append([None, m["meal"], i["name"], i["grams"], i["kcal"]])
                for cell in ws[ws.max_row]:
                    cell.border = THIN
    ws.append([])
    ws.append([f"Модель ИИ: {menu.get('model', '')}. Итоги дней посчитал бот."])


def _weights(wb: Workbook, rows) -> None:
    ws = _sheet(wb, "Дневник веса", ["Дата и время (МСК)", "Вес, кг", "Изменение от старта, кг"],
                [22, 10, 24])
    if not rows:
        ws.append(["Записей пока нет — нажмите в боте «⚖️ Дневник веса»."])
        return
    start = rows[0]["weight"]
    for r in rows:
        ws.append([r["created_at"], r["weight"], round(r["weight"] - start, 1)])
    if len(rows) >= 2:
        chart = LineChart()
        chart.title, chart.height, chart.width = "Динамика веса, кг", 8, 16
        chart.varyColors = False
        chart.legend = None
        chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl 3.1 по умолчанию прячет оси
        ws_min, ws_max = min(r["weight"] for r in rows), max(r["weight"] for r in rows)
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = int(ws_min) - 1, int(ws_max) + 2
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        chart.series[0].graphicalProperties.line.solidFill = GREEN
        chart.series[0].graphicalProperties.line.width = 28000
        chart.series[0].smooth = False
        chart.series[0].marker.symbol = "circle"
        chart.series[0].marker.graphicalProperties.solidFill = GREEN
        chart.series[0].marker.graphicalProperties.line.solidFill = GREEN
        ws.add_chart(chart, "E2")


def _favorites(wb: Workbook, rows) -> None:
    ws = _sheet(wb, "Мои рецепты", ["Рецепт", "Оценка", "Звёзды", "Добавлен (UTC)"], [44, 9, 10, 20])
    if not rows:
        ws.append(["Пока пусто — оцените рецепт звёздами в «🔍 Поиск рецептов»."])
        return
    for f in rows:
        ws.append([f["title"], f["rating"], "★" * f["rating"], f["added_at"]])


def build(user_id: int) -> bytes | None:
    """Книга .xlsx для пользователя или None, если анкеты нет."""
    p = storage.get_profile(user_id)
    if not p:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    _profile(wb, p)
    _menu(wb, storage.get_menu(user_id))
    _weights(wb, storage.weights(user_id))
    _favorites(wb, storage.favorites(user_id))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
