"""Проверка без Telegram и ИИ: дневник веса + отчёт Excel на временной базе.

  python tools/excel_check.py [out.xlsx]  — сохранит книгу для просмотра (по умолчанию не сохраняет)
"""
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook  # noqa: E402
from io import BytesIO  # noqa: E402

import storage  # noqa: E402

storage.DB = Path(tempfile.mkdtemp()) / "test.db"
import bot  # noqa: E402
import export  # noqa: E402
from nutrition import calc  # noqa: E402

UID = 1
P = {"sex": "m", "age": 40, "weight": 85.0, "height": 180, "activity": "mid", "goal": "lose"}
storage.init()
assert export.build(UID) is None, "без анкеты отчёта быть не должно"

storage.save_profile(UID, name="Тест", **P, **calc(**P))
for w in (85.0, 84.6, 84.1, 83.9):
    storage.add_weight(UID, w)
day = {"day": "Понедельник", "kcal": 1850, "protein": 120, "fat": 60, "carbs": 200, "meals": [
    {"meal": "Завтрак", "kcal": 450, "items": [{"name": "Овсянка на молоке", "grams": 250, "kcal": 300},
                                              {"name": "Яблоко", "grams": 150, "kcal": 150}]},
    {"meal": "Обед", "kcal": 700, "items": [{"name": "Гречка с курицей", "grams": 350, "kcal": 700}]}]}
storage.save_menu(UID, {"days": [day, day | {"day": "Вторник"}], "model": "openai/gpt-4o-mini", "target": 2000})
storage.rate(UID, "52795", "Курица по-индийски", 5)

print(re.sub(r"<[^>]+>", "", bot.weight_summary(storage.weights(UID))))
data = export.build(UID)
wb = load_workbook(BytesIO(data))
print("\nЛисты:", wb.sheetnames, f"· {len(data)} байт")
for ws in wb:
    print(f"— {ws.title}: {ws.max_row - 1} строк, графиков: {len(ws._charts)}")
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        print("   ", row)
assert wb["Дневник веса"]["C5"].value == -1.1 and len(wb["Дневник веса"]._charts) == 1
if len(sys.argv) > 1:
    Path(sys.argv[1]).write_bytes(data)
    print("\nСохранено:", sys.argv[1])
print("\nOK")
